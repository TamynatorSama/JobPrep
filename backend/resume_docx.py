"""Tailored-resume .docx renderer (split out of routes/application.py).

Kept in its own module so python-docx (~2s cold import) loads only when a
resume is actually built, not at sidecar boot. Pure rendering: plain-text
resume in, .docx bytes out.
"""
from __future__ import annotations

import io

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def _add_bottom_border(paragraph) -> None:
    """Attach a full-width horizontal rule under a paragraph (section divider).

    `w:sz` is in eighths-of-a-point — 8 = 1pt. `w:space` (in points) pushes
    the rule a few pt below the text so the heading doesn't sit on the line.
    """
    pPr = paragraph._p.get_or_add_pPr()
    # Drop any existing pBdr so we don't stack borders.
    for existing in pPr.findall(qn("w:pBdr")):
        pPr.remove(existing)
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "8")        # 1pt — visible without screaming
    bottom.set(qn("w:space"), "4")        # 4pt gap between text and rule
    bottom.set(qn("w:color"), "606060")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_run_with_inline_bold(paragraph, text: str) -> None:
    """Add runs to `paragraph` splitting `text` on '**...**' bold spans."""
    parts = text.split("**")
    for i, seg in enumerate(parts):
        if not seg:
            continue
        r = paragraph.add_run(seg)
        r.bold = (i % 2 == 1)


FONT_NAME = "Times New Roman"


def _set_font(run, *, size_pt: float | None = None, bold: bool | None = None,
              color: "RGBColor | None" = None) -> None:
    """Apply Times New Roman + optional size/bold/color to a run.

    python-docx does not always propagate `font.name` from the Normal style
    into runs with explicit rPr children, so we set it on every run we
    build to guarantee the whole resume renders in one font.
    """
    run.font.name = FONT_NAME
    # rFonts override — pins eastAsia + cs slots so Word doesn't substitute
    # back to Calibri for any character class.
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT_NAME)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def build_docx(resume_text: str) -> bytes:
    """Render the LLM's plain-text resume into a clean, ATS-friendly .docx.

    Layout rules:
        Line 1            → candidate name (centered, 20pt, bold)
        Line 2            → contact info (centered, 10pt, grey)
        '## SECTION'      → uppercase heading, 11.5pt bold, bottom-rule
        '### Subhead'     → 11pt bold (role / project / school)
        '- bullet'        → List Bullet style with hanging indent
        '**text**'        → inline bold
        '' (blank)        → tight spacer

    Times New Roman 11pt + 0.6in margins is the most ATS-portable combo.
    """
    doc = Document()

    # Tighter margins: 0.6in left/right, 0.5in top/bottom.
    section = doc.sections[0]
    section.top_margin    = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin   = Inches(0.6)
    section.right_margin  = Inches(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after  = Pt(2)
    normal.paragraph_format.line_spacing = 1.15

    lines = [ln.rstrip() for ln in resume_text.split("\n")]
    saw_name = False
    saw_contact = False

    GREY = RGBColor(0x55, 0x55, 0x55)
    BLACK_INK = RGBColor(0x10, 0x10, 0x10)

    for raw in lines:
        line = raw.strip()

        if not line:
            spacer = doc.add_paragraph("")
            spacer.paragraph_format.space_after = Pt(2)
            continue

        # ── Header: name + contact (first two non-empty lines) ───────────
        if not saw_name:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after  = Pt(0)
            p.alignment = 1  # WD_ALIGN_PARAGRAPH.CENTER
            _set_font(p.add_run(line), size_pt=20, bold=True, color=BLACK_INK)
            saw_name = True
            continue

        if not saw_contact:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after  = Pt(4)
            p.alignment = 1
            _set_font(p.add_run(line), size_pt=10, color=GREY)
            saw_contact = True
            continue

        # ── Section heading (## or #) ────────────────────────────────────
        if line.startswith("## "):
            p = doc.add_paragraph()
            # Modest space above so sections feel separated without burning
            # vertical real estate; smaller gap below so the rule + body
            # read as one block.
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after  = Pt(3)
            _set_font(p.add_run(line[3:].strip().upper()),
                      size_pt=11.5, bold=True, color=BLACK_INK)
            _add_bottom_border(p)
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after  = Pt(2)
            _set_font(p.add_run(line[2:].strip()), size_pt=13, bold=True)
            continue

        # ── Sub-heading (### role/project/school) ─────────────────────────
        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after  = Pt(1)
            _add_run_with_inline_bold(p, line[4:].strip())
            # Force the whole sub-head bold regardless of inline markers.
            for r in p.runs:
                _set_font(r, size_pt=11, bold=True)
            continue

        # ── Bullet ────────────────────────────────────────────────────────
        if line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            _add_run_with_inline_bold(p, line[2:].strip())
            for r in p.runs:
                _set_font(r)
            continue

        # ── Bare bold line e.g. "**Education**" ───────────────────────────
        if line.startswith("**") and line.endswith("**") and len(line) > 4:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after  = Pt(1)
            _set_font(p.add_run(line[2:-2]), bold=True)
            continue

        # ── Regular paragraph (with optional inline **bold** spans) ──────
        p = doc.add_paragraph()
        _add_run_with_inline_bold(p, line)
        for r in p.runs:
            _set_font(r)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
